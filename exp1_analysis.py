#!/usr/bin/env python3
"""Experiment 1 report: the requested graphs and tables, in the requested folders.

Usage (from anywhere):
    python3 exp1_report.py results/exp1
    python3 exp1_report.py results/exp1 --combined-out results/exp1_combined --target 25
    python3 exp1_report.py results/exp1 --load 0.8 --d 1     # if a raw folder mixes settings

Reads    <root>/<traffic>/raw/L<load>_D<0|1>_exp1_<pct>pct_run<replicate>.csv
Writes   <root>/<traffic>/graphs/    the 7 graphs for ONE traffic type (cbr, poisson, poissonBurst), PNG + PDF
         <root>/<traffic>/tables/    analysis tables for that traffic type
         <combined-out>/graphs/      the same 7 graphs with all traffic types together
         <combined-out>/tables/      analysis tables for all traffic types together (+ .md / .xlsx)
         (default <combined-out> = a folder called exp1_combined next to <root>)

Graphs (file name prefix = graph number)
    01 VO throughput vs adoption      02 VO mean delay vs adoption     03 VO p95 delay vs adoption
    04 VO p99 delay vs adoption       05 VO PDR vs adoption
    06 throughput of all ACs (bar)    07 mean delay of all ACs (bar, log scale)

Method: mean +/- 95% CI (Student t) over replicates; one replicate = one (seed, run) pair, and the same
replicate has the same random traffic at every adoption level.  VO_ALL = all VO stations together.  Delay
percentiles cannot be combined across the two station groups, so the "all VO" p95/p99 markers exist only
at 0% and 100% adoption; the legacy / P-EDCA group lines cover the mixed points.  p95/p99 in the tables are
the MEAN OF PER-REPLICATE percentiles, not percentiles of pooled packets.  Delay statistics only count
delivered packets - read them together with the PDR graph.
"""
import argparse, glob, math, os, re, sys
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 10,
    "axes.titlesize": 12,
    "axes.labelsize": 10,
    "xtick.labelsize": 9,
    "ytick.labelsize": 9,
    "legend.fontsize": 9,
    "figure.titlesize": 14,
    "axes.spines.top": False,
    "axes.spines.right": False,
    "axes.linewidth": 0.8,
    "lines.linewidth": 2.0,
    "lines.markersize": 6,
    "errorbar.capsize": 3,
    "grid.alpha": 0.22,
    "grid.linewidth": 0.6,
    "savefig.facecolor": "white",
    "figure.facecolor": "white",
})

# Denominator of the deliveredWithin10/20msRatio columns in the CSV: "tx" (delivered-in-time / transmitted)
# or "rx" (delivered-in-time / received).  Used as the weight when VO_LEGACY and VO_PEDCA are combined.
WITHIN_WEIGHT = "tx"

# ----------------------------------------------------------------------------- statistics
try:
    from scipy import stats as _st

    def tq(n):
        return float(_st.t.ppf(0.975, n - 1)) if n > 1 else float("nan")
except Exception:  # scipy not installed: Cornish-Fisher expansion of the t quantile (accurate for df >= 2)
    _T1, _T2 = 12.7062047, 4.30265273

    def tq(n):
        df = n - 1
        if df < 1:
            return float("nan")
        if df == 1:
            return _T1
        if df == 2:
            return _T2
        z = 1.959963985
        return z + (z**3 + z) / (4 * df) + (5 * z**5 + 16 * z**3 + 3 * z) / (96 * df**2) \
            + (3 * z**7 + 19 * z**5 + 17 * z**3 - 15 * z) / (384 * df**3)


def mci(x):
    """mean, 95% CI half-width, n (NaNs dropped)."""
    x = pd.Series(x, dtype=float).dropna()
    n = len(x)
    if n == 0:
        return np.nan, np.nan, 0
    if n == 1:
        return float(x.iloc[0]), np.nan, 1
    return float(x.mean()), tq(n) * float(x.std(ddof=1)) / math.sqrt(n), n


def fmt(mean, ci, dec=2, plus=False):
    if mean is None or (isinstance(mean, float) and np.isnan(mean)):
        return "n/a"
    s = f"{mean:+.{dec}f}" if plus else f"{mean:.{dec}f}"
    return s if (ci is None or np.isnan(ci)) else f"{s} ± {ci:.{dec}f}"


def verdict(mean, ci, target):
    if mean is None or np.isnan(mean):
        return "n/a"
    if not np.isnan(ci) and mean - ci >= target:
        return "MET"
    if mean >= target:
        return "likely (CI crosses target)"
    return "not met"


def nansum_or_nan(vals):
    v = np.array(vals, float)
    return float(np.nansum(v)) if (~np.isnan(v)).any() else np.nan


# ----------------------------------------------------------------------------- loading
FN = re.compile(r"L(?P<load>[\d.]+)_D(?P<d>[01])_exp1_(?P<pct>\d+)pct_run(?P<rep>\d+)\.csv$")


def load_traffic(raw_dir, load=None, dflag=None):
    frames = []
    for f in sorted(glob.glob(os.path.join(raw_dir, "*.csv"))):
        m = FN.search(os.path.basename(f))
        if not m or os.path.getsize(f) == 0:
            continue
        if load is not None and abs(float(m["load"]) - load) > 1e-9:
            continue
        if dflag is not None and int(m["d"]) != dflag:
            continue
        d = pd.read_csv(f)
        d["pct"] = int(m["pct"]); d["rep"] = int(m["rep"])
        d["load"] = float(m["load"]); d["dflag"] = int(m["d"])
        frames.append(d)
    if not frames:
        return None
    d = pd.concat(frames, ignore_index=True)
    combos = d[["load", "dflag"]].drop_duplicates()
    if len(combos) > 1:
        sys.exit(f"{raw_dir} mixes several (load, D) settings:\n{combos.to_string(index=False)}\n"
                 f"Separate them or select one with --load and --d.")
    n0 = len(d)
    d = d.drop_duplicates(subset=["pct", "rep", "ac"], keep="first")
    if len(d) < n0:
        print(f"  WARNING: dropped {n0 - len(d)} duplicate rows in {raw_dir}")
    return d


def ac_metrics(r):
    tx, rx = float(r["txPackets"]), float(r["rxPackets"])
    att, ret = float(r["macDataAttempts"]), float(r["macRetransmissions"])
    ds = float(r["dsCtsTransmissions"]) if "dsCtsTransmissions" in r else np.nan
    rts = float(r["stage2RtsTransmissions"]) if "stage2RtsTransmissions" in r else np.nan
    st = float(r["stationsInAc"])
    js = float(r["jitterSamples"]) if "jitterSamples" in r else rx
    return dict(stations=st, tx=tx, rx=rx, att=att, ret=ret, ds=ds, rts=rts, js=js,
                pdr=rx / tx if tx > 0 else np.nan, thr=float(r["throughputMbps"]),
                delay=float(r["meanDelayMs"]), p95=float(r["p95DelayMs"]), p99=float(r["p99DelayMs"]),
                jit=float(r["meanJitterMs"]),
                w10=float(r["deliveredWithin10msRatio"]) if "deliveredWithin10msRatio" in r else np.nan,
                w20=float(r["deliveredWithin20msRatio"]) if "deliveredWithin20msRatio" in r else np.nan)


def combine(ms):
    if len(ms) == 1:
        return dict(ms[0])
    tx = sum(m["tx"] for m in ms); rx = sum(m["rx"] for m in ms)

    def wavg(key, wkey):
        w = np.array([m[wkey] for m in ms], float); v = np.array([m[key] for m in ms], float)
        ok = ~np.isnan(v) & (w > 0)
        return float((v[ok] * w[ok]).sum() / w[ok].sum()) if ok.any() else np.nan

    return dict(stations=sum(m["stations"] for m in ms), tx=tx, rx=rx, att=sum(m["att"] for m in ms),
                ret=sum(m["ret"] for m in ms), ds=nansum_or_nan([m["ds"] for m in ms]),
                rts=nansum_or_nan([m["rts"] for m in ms]), js=sum(m["js"] for m in ms),
                pdr=rx / tx if tx > 0 else np.nan, thr=sum(m["thr"] for m in ms),
                delay=wavg("delay", "rx"), p95=np.nan, p99=np.nan, jit=wavg("jit", "js"),
                w10=wavg("w10", WITHIN_WEIGHT), w20=wavg("w20", WITHIN_WEIGHT))


def save_figure(fig, path, dpi=300):
    """Save publication-ready raster (PNG) and vector (PDF) versions."""
    base, _ = os.path.splitext(path)
    fig.savefig(base + ".png", dpi=dpi, bbox_inches="tight", facecolor="white")
    fig.savefig(base + ".pdf", bbox_inches="tight", facecolor="white")


def finish(m):
    return dict(pdr=m["pdr"], loss=1 - m["pdr"] if not np.isnan(m["pdr"]) else np.nan,
                meanDelayMs=m["delay"], p95DelayMs=m["p95"], p99DelayMs=m["p99"], meanJitterMs=m["jit"],
                throughputMbps=m["thr"], thrPerSta=m["thr"] / m["stations"] if m["stations"] else np.nan,
                retx=m["ret"] / m["att"] if m["att"] > 0 else np.nan,
                within10=m["w10"], within20=m["w20"], dsCts=m["ds"], stage2Rts=m["rts"], attempts=m["att"],
                dsPerAttempt=m["ds"] / m["att"] if m["att"] > 0 else np.nan,
                rtsPerDs=m["rts"] / m["ds"] if m["ds"] and m["ds"] > 0 else np.nan)


METRICS = ["pdr", "loss", "meanDelayMs", "p95DelayMs", "p99DelayMs", "meanJitterMs", "throughputMbps",
           "thrPerSta", "retx", "within10", "within20", "dsCts", "stage2Rts", "attempts", "dsPerAttempt", "rtsPerDs"]


def long_table(d):
    rows = []
    for (pct, rep), g in d.groupby(["pct", "rep"]):
        a = {r["ac"]: ac_metrics(r) for _, r in g.iterrows() if r["stationsInAc"] > 0}
        vo = [a[k] for k in ("AC_VO_LEGACY", "AC_VO_PEDCA") if k in a]
        if not vo:
            continue
        groups = {"VO_ALL": combine(vo)}
        if "AC_VO_LEGACY" in a: groups["VO_LEGACY"] = a["AC_VO_LEGACY"]
        if "AC_VO_PEDCA" in a: groups["VO_PEDCA"] = a["AC_VO_PEDCA"]
        for nm, ac in (("VI", "AC_VI"), ("BE", "AC_BE"), ("BK", "AC_BK")):
            if ac in a: groups[nm] = a[ac]
        total = sum(groups[k]["thr"] for k in ("VO_ALL", "VI", "BE", "BK") if k in groups)
        for gname, m in groups.items():
            rows.append(dict(group=gname, pct=pct, rep=rep, **finish(m)))
        tot = {k: np.nan for k in METRICS}
        tot["throughputMbps"] = total
        rows.append(dict(group="TOTAL", pct=pct, rep=rep, **tot))
    return pd.DataFrame(rows)


def summarise(L):
    recs = {}
    for (grp, pct), g in L.groupby(["group", "pct"]):
        row = {"n": g.rep.nunique()}
        for m in METRICS:
            row[m + "_mean"], row[m + "_ci"], row[m + "_n"] = mci(g[m])
        recs[(grp, pct)] = row
    S = pd.DataFrame.from_dict(recs, orient="index")
    S.index = pd.MultiIndex.from_tuples(S.index, names=["group", "pct"])
    return S


def get(S, grp, pct, m):
    try:
        r = S.loc[(grp, pct)]
        return float(r[m + "_mean"]), float(r[m + "_ci"])
    except KeyError:
        return np.nan, np.nan


def relative(L):
    """Per-replicate change versus the 0 % baseline (VO_ALL at 0 %)."""
    key = {(r.group, r.pct, r.rep): r for r in L.itertuples(index=False)}
    reps = sorted(L.rep.unique()); pcts = sorted(L.pct.unique()); rows = []

    def ok(b):
        return b is not None and not np.isnan(b) and b > 0

    def red(b, g):  # reduction in %
        return 100 * (b - g) / b if ok(b) and not np.isnan(g) else np.nan

    def chg(b, g):  # change in %
        return 100 * (g - b) / b if ok(b) and not np.isnan(g) else np.nan

    for rep in reps:
        b = key.get(("VO_ALL", 0, rep))
        if b is None:
            continue
        for pct in pcts:
            if pct == 0:
                continue
            for name, grp in (("Adopters (P-EDCA STAs)", "VO_PEDCA"), ("Non-adopters (legacy STAs)", "VO_LEGACY"),
                              ("System (all VO)", "VO_ALL")):
                g = key.get((grp, pct, rep))
                if g is None:
                    continue
                rows.append(dict(group=name, pct=pct, rep=rep,
                                 p95_red=red(b.p95DelayMs, g.p95DelayMs), p99_red=red(b.p99DelayMs, g.p99DelayMs),
                                 mean_red=red(b.meanDelayMs, g.meanDelayMs),
                                 loss_red=red(b.loss, g.loss) if b.loss >= 0.002 else np.nan,
                                 thr_gain=chg(b.thrPerSta, g.thrPerSta),
                                 w20_gain=100 * (g.within20 - b.within20) if not np.isnan(g.within20) else np.nan))
            for nm in ("VI", "BE", "BK", "TOTAL"):
                g, b0 = key.get((nm, pct, rep)), key.get((nm, 0, rep))
                if g is None or b0 is None:
                    continue
                rows.append(dict(group="cost:" + nm, pct=pct, rep=rep,
                                 thr_gain=chg(b0.throughputMbps, g.throughputMbps),
                                 pdr_chg=100 * (g.pdr - b0.pdr) if not np.isnan(g.pdr) else np.nan,
                                 delay_chg=chg(b0.meanDelayMs, g.meanDelayMs)))
    return pd.DataFrame(rows)


def rel_summary(R):
    recs = []
    for (grp, pct), g in R.groupby(["group", "pct"]):
        row = dict(group=grp, pct=pct, n=g.rep.nunique())
        for m in ("p95_red", "p99_red", "mean_red", "loss_red", "thr_gain", "w20_gain", "pdr_chg", "delay_chg"):
            if m in g.columns:
                row[m + "_mean"], row[m + "_ci"], row[m + "_n"] = mci(g[m])
        recs.append(row)
    return pd.DataFrame(recs)


# ----------------------------------------------------------------------------- tables
def build_tables(S, RS, target):
    T = {}
    pcts = sorted(S.index.get_level_values("pct").unique())
    # T1 system-level VO
    rows = []
    for p in pcts:
        r = {"P-EDCA adoption (%)": p, "replicates": int(S.loc[("VO_ALL", p)]["n"])}
        for lab, m, dec in (("VO PDR", "pdr", 4), ("loss (%)", "loss", 2), ("mean delay (ms)", "meanDelayMs", 2),
                            ("p95 delay (ms)", "p95DelayMs", 1), ("p99 delay (ms)", "p99DelayMs", 1),
                            ("jitter (ms)", "meanJitterMs", 2), ("VO throughput (Mbps)", "throughputMbps", 2),
                            ("retx ratio", "retx", 3), ("delivered <=10 ms", "within10", 3),
                            ("delivered <=20 ms", "within20", 3)):
            mean, ci = get(S, "VO_ALL", p, m)
            if m == "loss":
                mean, ci = mean * 100, ci * 100
            if np.isnan(mean) and m in ("p95DelayMs", "p99DelayMs"):
                r[lab] = "see group table"
            elif np.isnan(mean) and m in ("within10", "within20"):
                continue
            else:
                r[lab] = fmt(mean, ci, dec)
        rows.append(r)
    T["t1_vo_system"] = pd.DataFrame(rows)
    # T2 group comparison
    rows = []
    for p in pcts:
        for grp, lab in (("VO_LEGACY", "legacy (EDCA)"), ("VO_PEDCA", "P-EDCA")):
            if (grp, p) not in S.index:
                continue
            r = {"P-EDCA adoption (%)": p, "group": lab}
            for l, m, dec in (("PDR", "pdr", 4), ("loss (%)", "loss", 2), ("mean delay (ms)", "meanDelayMs", 2),
                              ("p95 (ms)", "p95DelayMs", 1), ("p99 (ms)", "p99DelayMs", 1),
                              ("jitter (ms)", "meanJitterMs", 2), ("retx ratio", "retx", 3),
                              ("delivered <=20 ms", "within20", 3)):
                mean, ci = get(S, grp, p, m)
                if m == "loss":
                    mean, ci = mean * 100, ci * 100
                if np.isnan(mean) and m == "within20":
                    continue
                r[l] = fmt(mean, ci, dec)
            rows.append(r)
    T["t2_vo_groups"] = pd.DataFrame(rows)
    # T3 cost to the other ACs
    rows = []
    for p in pcts:
        r = {"P-EDCA adoption (%)": p}
        for lab, grp, m, dec in (("VI thr (Mbps)", "VI", "throughputMbps", 2), ("VI PDR", "VI", "pdr", 4),
                                 ("VI mean delay (ms)", "VI", "meanDelayMs", 1), ("VI p95 (ms)", "VI", "p95DelayMs", 1),
                                 ("BE thr (Mbps)", "BE", "throughputMbps", 2), ("BE PDR", "BE", "pdr", 3),
                                 ("BK thr (Mbps)", "BK", "throughputMbps", 2), ("BK PDR", "BK", "pdr", 3),
                                 ("TOTAL thr (Mbps)", "TOTAL", "throughputMbps", 2)):
            mean, ci = get(S, grp, p, m)
            r[lab] = fmt(mean, ci, dec)
        rows.append(r)
    T["t3_other_acs"] = pd.DataFrame(rows)
    # T4 overhead
    rows = []
    for p in pcts:
        if ("VO_PEDCA", p) not in S.index:
            continue
        r = {"P-EDCA adoption (%)": p}
        for lab, m, dec in (("DS-CTS (count)", "dsCts", 0), ("Stage-2 RTS (count)", "stage2Rts", 0),
                            ("DS-CTS per data attempt", "dsPerAttempt", 3), ("Stage-2 RTS per DS-CTS", "rtsPerDs", 3)):
            mean, ci = get(S, "VO_PEDCA", p, m)
            r[lab] = fmt(mean, ci, dec)
        rows.append(r)
    T["t4_pedca_overhead"] = pd.DataFrame(rows)
    # T5 relative to baseline
    rows = []
    for _, r0 in RS.sort_values(["pct", "group"]).iterrows():
        if str(r0["group"]).startswith("cost:"):
            continue
        r = {"P-EDCA adoption (%)": int(r0["pct"]), "group": r0["group"], "replicates": int(r0["n"])}
        for lab, m in (("p95 delay reduction (%)", "p95_red"), ("p99 delay reduction (%)", "p99_red"),
                       ("mean delay reduction (%)", "mean_red"), ("dropped-packet reduction (%)", "loss_red"),
                       ("throughput per STA change (%)", "thr_gain")):
            mean, ci = r0.get(m + "_mean", np.nan), r0.get(m + "_ci", np.nan)
            r[lab] = fmt(mean, ci, 1, plus=True)
            if m == "loss_red":   # replicates with baseline loss < 0.2 % are excluded: report how many were used
                r["dropped-packet reduction n"] = int(r0.get("loss_red_n", 0))
            if m in ("p95_red", "loss_red", "thr_gain"):
                r[lab.split(" (")[0] + " vs " + f"{target:.0f}%"] = verdict(mean, ci, target)
        rows.append(r)
    T["t5_relative_vs_baseline"] = pd.DataFrame(rows)
    rows = []
    for _, r0 in RS.sort_values(["pct", "group"]).iterrows():
        if not str(r0["group"]).startswith("cost:"):
            continue
        r = {"P-EDCA adoption (%)": int(r0["pct"]), "AC": str(r0["group"]).split(":")[1], "replicates": int(r0["n"])}
        r["throughput change (%)"] = fmt(r0.get("thr_gain_mean", np.nan), r0.get("thr_gain_ci", np.nan), 1, plus=True)
        r["PDR change (points)"] = fmt(r0.get("pdr_chg_mean", np.nan), r0.get("pdr_chg_ci", np.nan), 2, plus=True)
        rows.append(r)
    T["t6_cost_vs_baseline"] = pd.DataFrame(rows)
    return T


def write_md(path, sections):
    def md(df):
        cols = list(df.columns)
        out = ["| " + " | ".join(map(str, cols)) + " |", "|" + "|".join("---" for _ in cols) + "|"]
        for _, r in df.iterrows():
            out.append("| " + " | ".join(str(r[c]) for c in cols) + " |")
        return "\n".join(out)

    with open(path, "w", encoding="utf-8") as f:
        f.write("# Experiment 1 - summary tables\n\nMean ± 95% CI over replicates. "
                "Reductions are relative to the 0% adoption baseline, paired per replicate.\n")
        for title, df in sections:
            f.write(f"\n### {title}\n\n{md(prepare_table(df))}\n")


# ----------------------------------------------------------------------------- plotting
TRAFFIC_ORDER = ["cbr", "poisson", "poissonBurst"]
TCOL = {"cbr": "#1f77b4", "poisson": "#2ca02c", "poissonBurst": "#d62728"}
C_SYS, C_LEG, C_PED = "#222222", "#d62728", "#1f77b4"
PCT_COLORS = {0: "#bdbdbd", 25: "#9ecae1", 50: "#6baed6", 75: "#3182bd", 100: "#08306b"}
GROUP_LABEL = {"VO_ALL": "all VO (system)", "VO_LEGACY": "legacy STAs (EDCA)", "VO_PEDCA": "P-EDCA STAs"}

TABLE_COLUMN_NAMES = {
    "P-EDCA adoption (%)": "P-EDCA adoption (%)",
    "replicates": "Replicates (n)",
    "VO PDR": "VO packet delivery ratio (fraction)",
    "loss (%)": "Packet loss (%)",
    "mean delay (ms)": "Mean delay (ms)",
    "p95 delay (ms)": "95th-percentile delay (ms)",
    "p99 delay (ms)": "99th-percentile delay (ms)",
    "jitter (ms)": "Mean jitter (ms)",
    "VO throughput (Mbps)": "Aggregate VO throughput (Mbps)",
    "throughput change (%)": "Throughput change (%)",
    "PDR change (points)": "PDR change (percentage points)",
}


def prepare_table(df):
    """Standardize table headings without changing numerical results."""
    out = df.copy()
    out.columns = [TABLE_COLUMN_NAMES.get(str(c), str(c)) for c in out.columns]
    return out


def errv(ci):
    return np.where(np.isnan(ci), 0, ci)


def series(S, grp, m, pcts):
    xs, ys, es = [], [], []
    for p in pcts:
        mean, ci = get(S, grp, p, m)
        if not np.isnan(mean):
            xs.append(p); ys.append(mean); es.append(ci)
    return np.array(xs), np.array(ys), errv(np.array(es))


def finish_axes(ax, title, ylabel, pcts):
    if title:
        ax.set_title(title, pad=10)
    ax.set_ylabel(ylabel)
    ax.set_xlabel("P-EDCA adoption among VO stations (%)", labelpad=8)
    ax.set_xticks(pcts)
    ax.set_xlim(min(pcts) - 3, max(pcts) + 3)
    ax.grid(True, axis="y", which="major")
    ax.grid(False, axis="x")
    ax.tick_params(axis="both", direction="out", length=3)
    ax.set_axisbelow(True)


VO_GRAPHS = [
    ("01_vo_throughput", "VO throughput", "VO throughput (Mbps)", "throughputMbps"),
    ("02_vo_mean_delay", "VO mean delay", "mean delay (ms)", "meanDelayMs"),
    ("03_vo_p95_delay", "VO 95th-percentile delay", "p95 delay (ms)", "p95DelayMs"),
    ("04_vo_p99_delay", "VO 99th-percentile delay", "p99 delay (ms)", "p99DelayMs"),
    ("05_vo_pdr", "VO packet delivery ratio", "PDR", "pdr"),
]
PERCENTILES = ("p95DelayMs", "p99DelayMs")


def draw_groups(ax, S, pcts, m, color=None, label_prefix="", ms=6, lw=1.8, alpha=1.0):
    """all-VO / legacy / P-EDCA lines for one traffic type on one axis."""
    kw = dict(capsize=3, alpha=alpha, lw=lw, markersize=ms)
    x, y, e = series(S, "VO_ALL", m, pcts)
    ax.errorbar(x, y, e, color=color or C_SYS, marker="D", ls="none" if m in PERCENTILES else ":", mfc="white",
                label=label_prefix + GROUP_LABEL["VO_ALL"], **kw)
    x, y, e = series(S, "VO_LEGACY", m, pcts)
    ax.errorbar(x, y, e, color=color or C_LEG, marker="o", ls="--", label=label_prefix + GROUP_LABEL["VO_LEGACY"], **kw)
    x, y, e = series(S, "VO_PEDCA", m, pcts)
    ax.errorbar(x, y, e, color=color or C_PED, marker="s", ls="-", label=label_prefix + GROUP_LABEL["VO_PEDCA"], **kw)


def pdr_floor(Ss, pcts):
    vals = [get(S, g, p, "pdr")[0] for S in Ss for g in ("VO_ALL", "VO_LEGACY", "VO_PEDCA") for p in pcts]
    vals = [v for v in vals if not np.isnan(v)]
    return max(0, min(vals) - 0.01) if vals else 0


def graph_vo_single(S, pcts, traffic, stem, title, ylabel, m, path, dpi):
    if m == "throughputMbps":  # total (all VO) + per-station by group
        fig, axs = plt.subplots(1, 2, figsize=(13, 5))
        x, y, e = series(S, "VO_ALL", m, pcts)
        axs[0].errorbar(x, y, e, color=C_SYS, marker="D", capsize=3, lw=1.8, label=GROUP_LABEL["VO_ALL"])
        finish_axes(axs[0], "total VO throughput", "VO throughput (Mbps)", pcts); axs[0].legend(fontsize=9)
        for g, c, mk, ls in (("VO_LEGACY", C_LEG, "o", "--"), ("VO_PEDCA", C_PED, "s", "-")):
            x, y, e = series(S, g, "thrPerSta", pcts)
            axs[1].errorbar(x, y, e, color=c, marker=mk, ls=ls, capsize=3, lw=1.8, label=GROUP_LABEL[g])
        finish_axes(axs[1], "throughput per VO station", "Mbps per station", pcts); axs[1].legend(fontsize=9)
    else:
        fig, ax = plt.subplots(figsize=(7.5, 5.2))
        draw_groups(ax, S, pcts, m)
        finish_axes(ax, "", ylabel, pcts); ax.legend(fontsize=9)
        if m == "pdr":
            ax.set_ylim(pdr_floor([S], pcts), 1.002)
    fig.suptitle(f"{title} vs P-EDCA adoption - {traffic} traffic (mean ± 95% CI)", fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, .95])
    save_figure(fig, path, dpi)
    plt.close(fig)


def graph_vo_combined(SS, pcts, stem, title, ylabel, m, path, dpi):
    traffics = [t for t in TRAFFIC_ORDER if t in SS]
    if m == "throughputMbps":
        fig, axs = plt.subplots(1, 2, figsize=(14, 5.5))
        for t in traffics:
            x, y, e = series(SS[t], "VO_ALL", m, pcts)
            axs[0].errorbar(x, y, e, color=TCOL[t], marker="D", capsize=3, lw=1.8, label=t)
            for g, mk, ls in (("VO_LEGACY", "o", "--"), ("VO_PEDCA", "s", "-")):
                x, y, e = series(SS[t], g, "thrPerSta", pcts)
                axs[1].errorbar(x, y, e, color=TCOL[t], marker=mk, ls=ls, capsize=3, lw=1.6, markersize=5)
        finish_axes(axs[0], "total VO throughput (all VO stations)", "VO throughput (Mbps)", pcts)
        axs[0].legend(fontsize=9, title="traffic")
        finish_axes(axs[1], "throughput per VO station", "Mbps per station", pcts)
        axs[1].legend(handles=[Line2D([], [], color="k", ls="--", marker="o", label="legacy STAs"),
                               Line2D([], [], color="k", ls="-", marker="s", label="P-EDCA STAs")], fontsize=9)
    else:
        fig, ax = plt.subplots(figsize=(9.5, 6))
        for t in traffics:
            draw_groups(ax, SS[t], pcts, m, color=TCOL[t], ms=5, lw=1.5, alpha=.95)
        finish_axes(ax, "", ylabel, pcts)
        h = [Line2D([], [], color=TCOL[t], lw=3, label=t) for t in traffics]
        h += [Line2D([], [], color="k", ls=":" if m not in PERCENTILES else "none", marker="D", mfc="white",
                     label="all VO (system)"),
              Line2D([], [], color="k", ls="--", marker="o", label="legacy STAs"),
              Line2D([], [], color="k", ls="-", marker="s", label="P-EDCA STAs")]
        ax.legend(handles=h, fontsize=9, ncol=2)
        if m == "pdr":
            ax.set_ylim(pdr_floor(list(SS.values()), pcts), 1.002)
    fig.suptitle(f"{title} vs P-EDCA adoption - all traffic types (mean ± 95% CI)", fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, .95])
    save_figure(fig, path, dpi)
    plt.close(fig)


AC_GROUPS = [("VO_ALL", "VO"), ("VI", "VI"), ("BE", "BE"), ("BK", "BK")]


def draw_ac_bars(ax, S, pcts, m, log=False):
    w = 0.8 / len(pcts); xs = np.arange(len(AC_GROUPS))
    for i, p in enumerate(pcts):
        vals, errs = [], []
        for g, _ in AC_GROUPS:
            mean, ci = get(S, g, p, m)
            vals.append(mean); errs.append(0 if np.isnan(ci) else ci)
        pos = xs + (i - (len(pcts) - 1) / 2) * w
        ax.bar(pos, np.nan_to_num(vals, nan=0.0), w, yerr=errs, color=PCT_COLORS.get(p, "gray"), capsize=2,
               edgecolor="white", linewidth=.4, label=f"{p}% P-EDCA")
        for xx, v in zip(pos, vals):
            if np.isnan(v):
                ax.text(xx, 0.02, "–", ha="center", va="bottom", fontsize=7, transform=ax.get_xaxis_transform())
    ax.set_xticks(xs); ax.set_xticklabels([n for _, n in AC_GROUPS], fontsize=10)
    ax.grid(alpha=.3, axis="y"); ax.tick_params(labelsize=9)
    if log:
        ax.set_yscale("log")


def graph_ac_bar_single(S, pcts, traffic, title, ylabel, m, log, path, dpi):
    fig, ax = plt.subplots(figsize=(9, 5.5))
    draw_ac_bars(ax, S, pcts, m, log)
    ax.set_ylabel(ylabel + (" (log scale)" if log else ""), fontsize=9); ax.set_xlabel("access category", fontsize=9)
    ax.legend(fontsize=8, title="adoption of P-EDCA among VO STAs", ncol=2)
    fig.suptitle(f"{title} of all access categories - {traffic} traffic (mean ± 95% CI)", fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, .95])
    save_figure(fig, path, dpi)
    plt.close(fig)


def graph_ac_bar_combined(SS, pcts, title, ylabel, m, log, path, dpi):
    traffics = [t for t in TRAFFIC_ORDER if t in SS]
    fig, axs = plt.subplots(1, len(traffics), figsize=(6.2 * len(traffics), 5.5), sharey=True, squeeze=False)
    for ax, t in zip(axs[0], traffics):
        draw_ac_bars(ax, SS[t], pcts, m, log)
        ax.set_title(t, fontsize=11); ax.set_xlabel("access category", fontsize=9)
    axs[0, 0].set_ylabel(ylabel + (" (log scale)" if log else ""), fontsize=9)
    axs[0, 0].legend(fontsize=8, title="adoption of P-EDCA among VO STAs", ncol=2)
    fig.suptitle(f"{title} of all access categories - all traffic types (mean ± 95% CI)", fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, .94])
    save_figure(fig, path, dpi)
    plt.close(fig)


# ----------------------------------------------------------------------------- main
def concat_tables(TT, name):
    parts = []
    for t in TRAFFIC_ORDER:
        if t in TT and name in TT[t]:
            x = TT[t][name].copy(); x.insert(0, "traffic", t); parts.append(x)
    return pd.concat(parts, ignore_index=True) if parts else None


def export_xlsx(path, sections):
    from openpyxl.styles import Font, Alignment, PatternFill
    used = set()
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for title, df in sections:
            sheet = re.sub(r"[^A-Za-z0-9_]", "_", title)[:31]
            base, k = sheet, 1
            while sheet in used:                      # keep sheet names unique after truncation
                k += 1; sheet = f"{base[:28]}_{k}"
            used.add(sheet)
            prepare_table(df).to_excel(xw, sheet_name=sheet, index=False, freeze_panes=(1, 1))
            ws = xw.sheets[sheet]
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="305496")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.row_dimensions[1].height = 32
            for col in ws.columns:
                w = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 12), 34)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("root", nargs="?", default=".", help="experiment folder holding <traffic>/raw (default: current folder)")
    ap.add_argument("--combined-out", default=None, help="folder for combined graphs and tables (default: <root>/../exp1_combined)")
    ap.add_argument("--target", type=float, default=25.0, help="relative-improvement target in %% for the verdict columns")
    ap.add_argument("--traffic", default=None, help='subset, e.g. "cbr poissonBurst"')
    ap.add_argument("--load", type=float, default=None, help="only use files with this load (L<load>)")
    ap.add_argument("--d", type=int, choices=[0, 1], default=None, help="only use files with this D flag")
    ap.add_argument("--dpi", type=int, default=300)
    a = ap.parse_args()
    root = os.path.abspath(a.root)
    comb = os.path.abspath(a.combined_out) if a.combined_out else os.path.join(os.path.dirname(root), "exp1_combined")
    wanted = a.traffic.split() if a.traffic else TRAFFIC_ORDER
    found = [t for t in wanted if os.path.isdir(os.path.join(root, t, "raw"))]
    if not found:
        sys.exit(f"No <traffic>/raw folders found under {root}")

    SS, TT = {}, {}
    for t in found:
        print(f"== {t}")
        d = load_traffic(os.path.join(root, t, "raw"), a.load, a.d)
        if d is None:
            print("  no usable CSV files, skipped"); continue
        L = long_table(d); S = summarise(L); R = relative(L)
        RS = rel_summary(R) if len(R) else pd.DataFrame()
        pcts = [int(p) for p in sorted(L.pct.unique())]
        if 0 not in pcts:
            print("  WARNING: no 0% adoption baseline found; relative tables will be empty")
        nrep = L.groupby("pct").rep.nunique()
        print(f"  adoption levels {pcts}; replicates per level {nrep.min()}-{nrep.max()}"
              + ("  (n<3: confidence intervals unreliable)" if nrep.min() < 3 else ""))
        gdir = os.path.join(root, t, "graphs"); tdir = os.path.join(root, t, "tables")
        os.makedirs(gdir, exist_ok=True); os.makedirs(tdir, exist_ok=True)
        for stem, title, ylabel, m in VO_GRAPHS:
            graph_vo_single(S, pcts, t, stem, title, ylabel, m, os.path.join(gdir, stem + ".png"), a.dpi)
        graph_ac_bar_single(S, pcts, t, "Throughput", "throughput (Mbps)", "throughputMbps", False,
                            os.path.join(gdir, "06_all_ac_throughput_bar.png"), a.dpi)
        graph_ac_bar_single(S, pcts, t, "Mean delay", "mean delay (ms)", "meanDelayMs", True,
                            os.path.join(gdir, "07_all_ac_mean_delay_bar.png"), a.dpi)
        T = build_tables(S, RS, a.target) if len(RS) else {}
        for name, df in T.items():
            prepare_table(df).to_csv(os.path.join(tdir, name + ".csv"), index=False, encoding="utf-8-sig")
        S.reset_index().to_csv(os.path.join(tdir, "numeric_summary_all_metrics.csv"), index=False)
        L.to_csv(os.path.join(tdir, "per_replicate_metrics.csv"), index=False)
        if len(RS):
            RS.to_csv(os.path.join(tdir, "numeric_relative_vs_baseline.csv"), index=False)
        SS[t], TT[t] = S, T
        print(f"  graphs -> {gdir}\n  tables -> {tdir}")

    if not SS:
        sys.exit("Nothing to report.")
    cg = os.path.join(comb, "graphs"); ct = os.path.join(comb, "tables")
    os.makedirs(cg, exist_ok=True); os.makedirs(ct, exist_ok=True)
    pcts = sorted({int(p) for S in SS.values() for p in S.index.get_level_values("pct")})
    for stem, title, ylabel, m in VO_GRAPHS:
        graph_vo_combined(SS, pcts, stem, title, ylabel, m, os.path.join(cg, stem + "_combined.png"), a.dpi)
    graph_ac_bar_combined(SS, pcts, "Throughput", "throughput (Mbps)", "throughputMbps", False,
                          os.path.join(cg, "06_all_ac_throughput_bar_combined.png"), a.dpi)
    graph_ac_bar_combined(SS, pcts, "Mean delay", "mean delay (ms)", "meanDelayMs", True,
                          os.path.join(cg, "07_all_ac_mean_delay_bar_combined.png"), a.dpi)

    titles = {"t1_vo_system": "VO, all stations together", "t2_vo_groups": "VO, legacy vs P-EDCA stations",
              "t3_other_acs": "VI / BE / BK and total throughput", "t4_pedca_overhead": "P-EDCA overhead",
              "t5_relative_vs_baseline": f"Change vs 0% adoption baseline ({a.target:.0f}% target verdicts)",
              "t6_cost_vs_baseline": "Cost to VI / BE / BK / total vs baseline"}
    sections = []
    for name, ttl in titles.items():
        df = concat_tables(TT, name)
        if df is not None:
            prepare_table(df).to_csv(os.path.join(ct, f"all_traffic_{name}.csv"), index=False, encoding="utf-8-sig")
            sections.append((ttl + " - all traffic types", df))
    write_md(os.path.join(comb, "exp1_tables.md"), sections)
    try:
        export_xlsx(os.path.join(comb, "exp1_tables.xlsx"), sections)
    except Exception as e:
        print(f"Excel export skipped: {e}")


if __name__ == "__main__":
    main()
